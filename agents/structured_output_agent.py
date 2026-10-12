"""
Agent 9 — Structured Output Agent
Generates JSON, CSV, Excel exports with provenance and validation details.
"""
from __future__ import annotations

import csv
import io
import json
from datetime import datetime
from typing import Any, Optional

from core.logging_config import get_logger

logger = get_logger(__name__)


def _build_export_record(doc: dict, extracted_fields: list[dict],
                          validation_findings: list[dict],
                          review_decisions: list[dict],
                          risk_result: dict) -> dict:
    """Build a complete export record for one document."""
    # Merge extracted fields, applying reviewer corrections
    field_values: dict[str, Any] = {}
    for f in extracted_fields:
        fname = f.get("field_name", "")
        field_values[fname] = {
            "value": f.get("normalized_value"),
            "source": f.get("extraction_method", "text"),
            "confidence": f.get("confidence", 1.0),
            "page": f.get("page_number"),
        }

    # Apply corrections
    corrections: dict[str, Any] = {}
    for rd in review_decisions:
        fname = rd.get("field_name")
        if fname and rd.get("action") == "correct":
            corrections[fname] = rd.get("corrected_value")
            if fname in field_values:
                field_values[fname]["corrected_value"] = rd.get("corrected_value")
                field_values[fname]["reviewer"] = rd.get("reviewer")
                field_values[fname]["reviewed_at"] = rd.get("reviewed_at")

    return {
        "schema_version": "1.0",
        "doc_id": doc.get("doc_id"),
        "filename": doc.get("filename"),
        "doc_type": doc.get("doc_type"),
        "status": doc.get("status"),
        "risk_score": doc.get("risk_score"),
        "risk_label": doc.get("risk_label"),
        "processing_mode": doc.get("processing_mode", "local"),
        "language": doc.get("language", "en"),
        "classification_confidence": doc.get("classification_confidence"),
        "extracted_fields": field_values,
        "reviewer_corrections": corrections,
        "validation_findings": [
            {k: v for k, v in f.items() if k not in ("id",)}
            for f in validation_findings
        ],
        "exported_at": datetime.utcnow().isoformat(),
    }


def export_json(doc: dict, extracted_fields: list[dict],
                validation_findings: list[dict],
                review_decisions: list[dict],
                risk_result: dict) -> bytes:
    """Returns JSON bytes."""
    record = _build_export_record(doc, extracted_fields, validation_findings,
                                   review_decisions, risk_result)
    return json.dumps(record, indent=2, ensure_ascii=False).encode("utf-8")


def export_csv(doc: dict, extracted_fields: list[dict],
               validation_findings: list[dict],
               review_decisions: list[dict],
               risk_result: dict) -> bytes:
    """Returns CSV bytes (field-level rows)."""
    record = _build_export_record(doc, extracted_fields, validation_findings,
                                   review_decisions, risk_result)
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["doc_id", "doc_type", "status", "risk_score",
                     "field_name", "value", "corrected_value", "source",
                     "confidence", "page"])
    for fname, fdata in record["extracted_fields"].items():
        writer.writerow([
            record["doc_id"], record["doc_type"], record["status"],
            record["risk_score"], fname,
            fdata.get("value"), fdata.get("corrected_value"),
            fdata.get("source"), fdata.get("confidence"), fdata.get("page"),
        ])
    return buf.getvalue().encode("utf-8")


def export_excel(doc: dict, extracted_fields: list[dict],
                 validation_findings: list[dict],
                 review_decisions: list[dict],
                 risk_result: dict) -> Optional[bytes]:
    """Returns Excel bytes, or None if openpyxl unavailable."""
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment
    except ImportError:
        logger.warning("openpyxl not available — Excel export unavailable.")
        return None

    record = _build_export_record(doc, extracted_fields, validation_findings,
                                   review_decisions, risk_result)
    wb = openpyxl.Workbook()

    # Sheet 1: Summary
    ws1 = wb.active
    ws1.title = "Summary"
    header_fill = PatternFill("solid", fgColor="1F3864")
    header_font = Font(bold=True, color="FFFFFF")
    summary_data = [
        ("Document ID", record["doc_id"]),
        ("Filename", record["filename"]),
        ("Document Type", record["doc_type"]),
        ("Status", record["status"]),
        ("Risk Score", record["risk_score"]),
        ("Risk Label", record["risk_label"]),
        ("Processing Mode", record["processing_mode"]),
        ("Exported At", record["exported_at"]),
    ]
    for r, (k, v) in enumerate(summary_data, 1):
        ws1.cell(r, 1, k).font = Font(bold=True)
        ws1.cell(r, 2, str(v) if v is not None else "")

    # Sheet 2: Fields
    ws2 = wb.create_sheet("Extracted Fields")
    headers = ["Field", "Value", "Corrected Value", "Source", "Confidence", "Page"]
    for col, h in enumerate(headers, 1):
        cell = ws2.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
    for row, (fname, fdata) in enumerate(record["extracted_fields"].items(), 2):
        ws2.cell(row, 1, fname)
        ws2.cell(row, 2, str(fdata.get("value") or ""))
        ws2.cell(row, 3, str(fdata.get("corrected_value") or ""))
        ws2.cell(row, 4, fdata.get("source", ""))
        ws2.cell(row, 5, fdata.get("confidence", ""))
        ws2.cell(row, 6, fdata.get("page") or "")

    # Sheet 3: Findings
    ws3 = wb.create_sheet("Validation Findings")
    find_headers = ["Rule ID", "Field", "Severity", "Message", "Expected", "Observed", "Passed"]
    for col, h in enumerate(find_headers, 1):
        cell = ws3.cell(1, col, h)
        cell.fill = header_fill
        cell.font = header_font
    severity_colors = {"Critical": "FF0000", "High": "FF6600", "Medium": "FFAA00",
                       "Low": "FFFF00", "Info": "CCCCCC"}
    for row, f in enumerate(record["validation_findings"], 2):
        ws3.cell(row, 1, f.get("rule_id", ""))
        ws3.cell(row, 2, f.get("field") or "")
        sev = f.get("severity", "")
        sev_cell = ws3.cell(row, 3, sev)
        color = severity_colors.get(sev, "FFFFFF")
        sev_cell.fill = PatternFill("solid", fgColor=color)
        ws3.cell(row, 4, f.get("message", ""))
        ws3.cell(row, 5, f.get("expected") or "")
        ws3.cell(row, 6, f.get("observed") or "")
        ws3.cell(row, 7, "✓" if f.get("passed") else "✗")

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_batch_csv(documents: list[dict]) -> bytes:
    """Returns a summary CSV for multiple documents."""
    buf = io.StringIO()
    writer = csv.writer(buf)
    writer.writerow(["doc_id", "filename", "doc_type", "status", "risk_score",
                     "risk_label", "classification_confidence", "page_count",
                     "created_at"])
    for d in documents:
        writer.writerow([
            d.get("doc_id"), d.get("filename"), d.get("doc_type"),
            d.get("status"), d.get("risk_score"), d.get("risk_label"),
            d.get("classification_confidence"), d.get("page_count"),
            d.get("created_at"),
        ])
    return buf.getvalue().encode("utf-8")


def generate_summary_report(doc: dict, extracted_fields: list[dict],
                             validation_findings: list[dict]) -> str:
    """Generate a human-readable text summary."""
    failures = [f for f in validation_findings if not f.get("passed")]
    lines = [
        f"DocuMind AI — Processing Summary",
        f"{'='*50}",
        f"Document  : {doc.get('filename')}",
        f"Type      : {doc.get('doc_type')}",
        f"Status    : {doc.get('status')}",
        f"Risk Score: {doc.get('risk_score', 0):.1%}",
        f"",
        f"Extracted {len(extracted_fields)} fields",
        f"Validation: {len(failures)} issue(s) found",
        f"",
    ]
    if failures:
        lines.append("Issues Found:")
        for f in failures:
            sev = f.get("severity", "?")
            msg = f.get("message", "")
            lines.append(f"  [{sev}] {msg}")
    else:
        lines.append("No validation issues found.")
    lines.append(f"\nGenerated: {datetime.utcnow().isoformat()}Z")
    return "\n".join(lines)
